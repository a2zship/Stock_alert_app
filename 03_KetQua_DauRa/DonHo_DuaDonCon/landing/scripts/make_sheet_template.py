"""Tạo file Excel mẫu cho bảng dữ liệu (tải lên Google Drive, mở bằng Google Sheets, công bố CSV).
Chạy: python scripts/make_sheet_template.py
Sửa phần CẤU HÌNH bên dưới cho khớp src/content.ts của ngành (tên danh mục tiếng Việt, nhóm lọc, tình trạng, tên cột).
Tên cột phải nằm trong danh sách COLS của src/data/items.ts (so khớp không dấu) — nếu đặt tên mới thì thêm vào COLS.
"""
import sys
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.stdout.reconfigure(encoding='utf-8')  # console Windows mặc định cp1252, in tiếng Việt bị lỗi

# ======================= CẤU HÌNH THEO NGÀNH =======================
SHEET_NAME = 'DanhSachTruong'
CATEGORIES = ['Mầm non', 'Tiểu học', 'THCS']  # cột Cấp học
SEGMENTS = []  # cột Tỉnh/thành: gõ tay, web tự thêm vào bộ lọc
STATUS = ['Đang đón', 'Đang tìm người đón', 'Tạm ngừng']  # hiện / hiện kèm nhãn / ẩn khỏi web
# (tên cột, độ rộng). Cột Cấp học / Tình trạng có danh sách thả xuống.
COLUMNS = [
    ('Mã', 10), ('Tên trường', 40), ('Cấp học', 14), ('Tỉnh/thành', 18), ('Phường/xã', 22), ('Số lớp', 10),
    ('Tình trạng', 20), ('Ghi chú', 40), ('Ngày cập nhật', 16),
]
# ===================================================================

ROOT = Path(__file__).resolve().parent.parent
today = date.today()
OUT = ROOT / 'docs' / f'Mau_Bang_DuLieu_{today:%d-%m-%Y}.xlsx'
HEAD = [c for c, _ in COLUMNS]
col = lambda name: chr(ord('A') + HEAD.index(name))  # noqa: E731 (tối đa 26 cột)

wb = Workbook()
ws = wb.active
ws.title = SHEET_NAME
ws.append(HEAD)
sample = {h: '' for h in HEAD}
sample.update({'Mã': 'T001', 'Tên trường': '[Tên trường đầy đủ]', 'Cấp học': CATEGORIES[0], 'Tỉnh/thành': '[Tỉnh/thành]',
               'Phường/xã': '[Phường/xã]', 'Số lớp': 1, 'Tình trạng': STATUS[1], 'Ghi chú': 'XOÁ DÒNG MẪU NÀY',
               'Ngày cập nhật': f'{today:%d/%m/%Y}'})
ws.append([sample[h] for h in HEAD])

bold, fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1E1B4B')
for i, (_, w) in enumerate(COLUMNS, start=1):
    c = ws.cell(row=1, column=i)
    c.font, c.fill = bold, fill
    c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[1].height = 32
ws.freeze_panes = 'C2'
if False:  # bảng trường không có cột giá
    for row in ws.iter_rows(min_row=2, min_col=HEAD.index('Giá (đồng)') + 1, max_col=HEAD.index('Giá (đồng)') + 1):
        for c in row:
            c.number_format = '#,##0'


def drop(name, items, allow_other=False):
    if name not in HEAD:
        return
    dv = DataValidation(type='list', formula1='"' + ','.join(items) + '"', allow_blank=True, showErrorMessage=not allow_other)
    dv.error = 'Chọn một giá trị trong danh sách.'
    ws.add_data_validation(dv)
    dv.add(f'{col(name)}2:{col(name)}1000')


drop('Cấp học', CATEGORIES)
drop('Tình trạng', STATUS)

guide = wb.create_sheet('HuongDan')
for line in [
    'CÁCH CẬP NHẬT DANH SÁCH TRƯỜNG',
    f'1. Mỗi trường một dòng ở trang tính {SHEET_NAME}. Không đổi tên trang tính và dòng tiêu đề.',
    '2. Chỉ ghi trường đã được quản trị Đón Hộ gọi xác minh (trang Quản trị: phễu trường trong app).',
    '3. Tên trường ghi đầy đủ, đúng như trong app. Cấp học và Tình trạng: chọn trong danh sách thả xuống.',
    '4. Tỉnh/thành, Phường/xã: gõ đúng tên hành chính như trong app. Tỉnh mới tự hiện trong bộ lọc của web.',
    f'5. "{STATUS[0]}": trường đã có lượt đón. "{STATUS[1]}": đã mở trường, đang mời phụ huynh, người đón. "{STATUS[2]}": web ẩn trường.',
    '6. Số lớp: số lớp đang dùng Đón Hộ (xem cột Lớp ở trang phễu trường). Để trống nếu chưa rõ.',
    '7. Web tự đọc lại bảng khoảng 5 phút một lần.',
    '',
    'KHÔNG GHI VÀO BẢNG: số điện thoại, tên cô giáo, tên phụ huynh, tên bé, lớp của bé, số tài khoản. Bảng công bố thì ai có link đều đọc được.',
]:
    guide.append([line])
guide['A1'].font = Font(bold=True, size=13)
guide.cell(row=guide.max_row, column=1).font = Font(bold=True, color='9A2A12')
guide.column_dimensions['A'].width = 120

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(OUT)
